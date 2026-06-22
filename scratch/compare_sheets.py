import openpyxl
import os

orig_path = os.path.abspath(os.path.join(os.path.dirname(__file__), './orig_StageStepTable.xlsx'))
current_path = os.path.abspath(os.path.join(os.path.dirname(__file__), '../public/StageStepTable.xlsx'))

wb_orig = openpyxl.load_workbook(orig_path)
wb_curr = openpyxl.load_workbook(current_path)

sheet_orig = wb_orig.active
sheet_curr = wb_curr.active

print("Orig dimensions:", sheet_orig.dimensions)
print("Curr dimensions:", sheet_curr.dimensions)

diff_count = 0
max_diffs_to_show = 50

for r in range(1, max(sheet_orig.max_row, sheet_curr.max_row) + 1):
    for c in range(1, max(sheet_orig.max_column, sheet_curr.max_column) + 1):
        v_orig = sheet_orig.cell(row=r, column=c).value
        v_curr = sheet_curr.cell(row=r, column=c).value
        
        if v_orig != v_curr:
            diff_count += 1
            if diff_count <= max_diffs_to_show:
                print(f"Diff at Cell({r}, {c}) [{openpyxl.utils.get_column_letter(c)}{r}]:")
                print(f"  Orig: {v_orig} (type: {type(v_orig).__name__ if v_orig is not None else 'None'})")
                print(f"  Curr: {v_curr} (type: {type(v_curr).__name__ if v_curr is not None else 'None'})")

print(f"Total differences found: {diff_count}")
